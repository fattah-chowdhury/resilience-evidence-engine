"""Small bounded local ingestion; URLs remain metadata and are never crawled."""

import csv
import io
import json
import xml.etree.ElementTree as ET
import zipfile
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path

from pydantic import ValidationError

from ree.models import Record, digest


class PlainText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "div", "br", "li"}:
            self.parts.append(" ")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        self.parts.append(" ")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def plain_html(text):
    parser = PlainText()
    parser.feed(text)
    return " ".join("".join(parser.parts).split())


def raw_rows(path, max_bytes):
    if path.stat().st_size > max_bytes:
        raise ValueError("Input exceeds collection.max_bytes")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > max_bytes * 10:
                raise ValueError("Excel expanded size exceeds budget")
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("Excel input needs pip install '.[excel]'") from error
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            rows = book.active.iter_rows(values_only=True)
            headers = list(next(rows))
            if len(set(headers)) != len(headers) or any(not isinstance(h, str) for h in headers):
                raise ValueError("Excel headers must be distinct strings")
            for row in rows:
                yield dict(zip(headers, row))
        finally:
            book.close()
        return
    if suffix == ".parquet":
        try:
            import pyarrow.parquet as pq
        except ImportError as error:
            raise ValueError("Parquet input needs pip install '.[columnar]'") from error
        for batch in pq.ParquetFile(path).iter_batches(batch_size=256):
            yield from batch.to_pylist()
        return
    if suffix == ".gpkg":
        try:
            import geopandas as gpd
        except ImportError as error:
            raise ValueError("GeoPackage input needs pip install '.[gis]'") from error
        frame = gpd.read_file(path)
        if frame.crs is None:
            raise ValueError("GeoPackage must declare its CRS")
        frame = frame.to_crs(4326)
        yield from json.loads(frame.to_json())["features"]
        return
    text = path.read_text(encoding="utf-8-sig")
    if suffix in {".csv", ".tsv"}:
        reader = csv.DictReader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",")
        if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError("CSV requires distinct column headers")
        yield from reader
    elif suffix in {".json", ".geojson"}:
        value = json.loads(text)
        if isinstance(value, list):
            yield from value
        elif isinstance(value, dict) and value.get("type") == "FeatureCollection":
            if "crs" in value:
                raise ValueError("GeoJSON must be WGS84 without a legacy crs member")
            yield from value["features"]
        else:
            raise ValueError("JSON input must be a list or GeoJSON FeatureCollection")
    elif suffix in {".jsonl", ".ndjson"}:
        for line in text.splitlines():
            if line.strip():
                try:
                    yield json.loads(line)
                except json.JSONDecodeError:
                    yield {"__parse_error__": "Invalid JSON line"}
    elif suffix in {".txt", ".html", ".htm"}:
        yield {"text": plain_html(text) if suffix != ".txt" else text}
    elif suffix in {".rss", ".xml", ".atom"}:
        if "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
            raise ValueError("DTD and entity declarations are not allowed")
        root = ET.fromstring(text)
        items = root.findall(".//item") or root.findall("{http://www.w3.org/2005/Atom}entry")
        for item in items:
            fields = {child.tag.rsplit("}", 1)[-1]: child for child in item}
            def content(name, nodes=fields):
                node = nodes.get(name)
                return "" if node is None else "".join(node.itertext())
            link = fields.get("link")
            yield {"source_record_id": content("guid") or content("id") or digest(ET.tostring(item).hex()),
                   "title": content("title"),
                   "text": plain_html(content("description") or content("summary") or content("content")
                                      or content("title")),
                   "url": "" if link is None else (link.get("href") or link.text or "")}
    else:
        raise ValueError(f"Unsupported input extension: {suffix}")


def convert_row(raw, columns):
    if not isinstance(raw, dict):
        raise ValueError("Each input record must be an object")  # noqa: TRY004 - invalid input data
    if raw.get("type") == "Feature":
        raw = {**raw.get("properties", {}), "geometry": raw.get("geometry"),
               "source_record_id": str(raw.get("id") or digest(raw))}
    row = {columns.get(k, k): v for k, v in raw.items() if v is not None and v != ""}
    if len(row) < len([v for v in raw.values() if v is not None and v != ""]):
        raise ValueError("Column mapping creates duplicate target fields")
    for key in {"geometry", "quantities", "acquisition"} & row.keys():
        if isinstance(row[key], str):
            row[key] = json.loads(row[key])
    if "longitude" in row or "latitude" in row:
        if not {"longitude", "latitude"} <= row.keys():
            raise ValueError("Both longitude and latitude are required")
        if "geometry" in row:
            raise ValueError("Supply geometry or longitude/latitude, not both")
        row["geometry"] = {"type": "Point", "coordinates":
                           [float(row.pop("longitude")), float(row.pop("latitude"))]}
    if row.get("geometry") and "geometry_role" not in row:
        row["geometry_role"] = "unknown" if row["geometry"]["type"] == "Point" else "area"
    row.setdefault("source_record_id", digest(row))
    row["source_record_id"] = str(row["source_record_id"])
    return row


def ingest_file(spec, base, max_records, max_bytes):
    path = (base / spec.path).resolve()
    input_hash = sha256(path.read_bytes()).hexdigest() if path.stat().st_size <= max_bytes else None
    source_id = "local_" + digest({"name": Path(spec.path).name,
                                 "rights": spec.model_dump(exclude={"path", "columns"})})
    source = {"id": source_id, "provider": "User-supplied local evidence", "origin_group": None,
              "license": spec.license, "redistribution_allowed": spec.redistribution_allowed,
              "raw_content_storage_allowed": True, "status": "user_declared"}
    records, errors = [], []
    seen = 0
    for i, raw in enumerate(raw_rows(path, max_bytes), 1):
        if i > max_records:
            errors.append({"row": i, "reason": "record_budget_exceeded"})
            break
        seen += 1
        try:
            row = convert_row(raw, spec.columns)
            row["source_id"] = source_id
            sensitive = row.get("sensitive", False)
            if isinstance(sensitive, str):
                if sensitive.lower() not in {"true", "false", "1", "0"}:
                    raise ValueError("sensitive must be true or false")
                sensitive = sensitive.lower() in {"true", "1"}
            row["sensitive"] = spec.sensitive or sensitive
            row["acquisition"] = {"kind": "local_file", "input_hash": input_hash, "row": i,
                                  "declared_metadata": row.get("acquisition", {})}
            records.append(Record.model_validate(row))
        except (ValueError, TypeError, KeyError, ValidationError) as error:
            reason = type(error).__name__
            if isinstance(error, ValidationError):
                reason = "; ".join(".".join(map(str, x["loc"])) + ": " + x["msg"]
                                   for x in error.errors(include_input=False))
            errors.append({"row": i, "reason": reason})
    return records, source, {"hash": input_hash, "rows_seen": seen}, errors
